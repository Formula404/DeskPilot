from __future__ import annotations

import json
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient

from backend.app.agent import nodes
from backend.app.agent.intents import detect_intent
from backend.app.core.config import get_settings
from backend.app.db.connection import init_db
from backend.app.db.connection import connect
from backend.app.db.repository import set_setting
from backend.app.knowledge import compiler
from backend.app.knowledge.compiler import KnowledgeCompileError, compile_source, review_proposal
from backend.app.knowledge.lint import lint_knowledge
from backend.app.knowledge.indexer import rebuild_index
from backend.app.knowledge.markdown import parse_markdown_document, source_chunks
from backend.app.knowledge.chinese_search import search_terms
from backend.app.knowledge.obsidian import export_obsidian_vault
from backend.app.knowledge.paths import KNOWLEDGE_ROOT_KEY, ensure_knowledge_dirs, from_knowledge_relative
from backend.app.knowledge.profiles import create_profile, switch_profile
from backend.app.knowledge.repository import get_note, get_proposal, get_source, list_proposals
from backend.app.knowledge.retrieval import answer_knowledge, search_knowledge
from backend.app.knowledge.settings import get_knowledge_settings, save_knowledge_settings
from backend.app.knowledge.source_service import canonicalize_url, ingest_content, ingest_file
from backend.app.knowledge.storage import change_storage_directory
from backend.app.knowledge.paths import knowledge_root
from backend.app.knowledge.catalog import catalog_status, refresh_catalogs
from backend.app.knowledge.lint import semantic_lint_knowledge
from backend.app.knowledge.compiler import promote_answer
from backend.app.knowledge.web_monitor import check_web_source, get_watch, set_watch
from backend.app.main import app
from backend.app.schemas.common import ToolResult


@pytest.fixture
def knowledge_data_dir(tmp_path):
    settings = get_settings()
    original = settings.data_dir
    settings.data_dir = tmp_path / "data"
    init_db()
    ensure_knowledge_dirs()
    try:
        yield settings.data_dir
    finally:
        settings.data_dir = original


def test_canonicalize_url_removes_tracking_parameters() -> None:
    assert canonicalize_url("HTTPS://Example.com/post?utm_source=x&id=4#part") == "https://example.com/post?id=4"


def test_knowledge_intents_have_priority() -> None:
    assert detect_intent("把当前网页加入知识库") == "knowledge_ingest"
    assert detect_intent("在知识库里查询 LangGraph") == "knowledge_query"
    assert detect_intent("检查知识库有没有断链") == "knowledge_maintenance"
    assert detect_intent("总结知识库中的 Python 内容") == "knowledge_query"
    assert detect_intent("从知识库里导出表格到 Excel") == "web_table_export"
    assert detect_intent("同意提案 prop_abc123") == "knowledge_review"


@pytest.mark.asyncio
async def test_review_requires_explicit_decision(monkeypatch: pytest.MonkeyPatch) -> None:
    called = False

    async def fake_run(*args, **kwargs):
        nonlocal called
        called = True
        return args[0], ToolResult(ok=True)

    monkeypatch.setattr(nodes, "_run_knowledge_tool", fake_run)
    result = await nodes.review_knowledge(
        {"task_id": "task-review", "user_input": "审核提案 prop_abc123"}
    )
    assert not called
    assert "尚未处理" in result["final_response"]


@pytest.mark.asyncio
async def test_maintenance_check_does_not_rebuild(monkeypatch: pytest.MonkeyPatch) -> None:
    called_tool = ""

    async def fake_run(state, tool_name, payload):
        nonlocal called_tool
        called_tool = tool_name
        return state, ToolResult(
            ok=True,
            data={"summary": {"errors": 0, "warnings": 0}},
        )

    monkeypatch.setattr(nodes, "_run_knowledge_tool", fake_run)
    await nodes.maintain_knowledge(
        {"task_id": "task-maintain", "user_input": "检查知识库是否需要重建索引"}
    )
    assert called_tool == "knowledge.lint"


@pytest.mark.asyncio
async def test_rebuild_requires_confirmation(monkeypatch: pytest.MonkeyPatch) -> None:
    called = False

    async def fake_run(*args, **kwargs):
        nonlocal called
        called = True
        return args[0], ToolResult(ok=True)

    monkeypatch.setattr(nodes, "_run_knowledge_tool", fake_run)
    result = await nodes.maintain_knowledge(
        {"task_id": "task-maintain", "user_input": "重建知识索引"}
    )
    assert not called
    assert "确认重建知识索引" in result["final_response"]


@pytest.mark.asyncio
async def test_ingest_accepts_forward_slash_windows_path(monkeypatch: pytest.MonkeyPatch) -> None:
    captured: dict = {}

    async def fake_run(state, tool_name, payload):
        captured.update({"tool": tool_name, "payload": payload})
        return state, ToolResult(
            ok=True,
            data={"status": "created", "title": "notes", "compilation": None},
        )

    monkeypatch.setattr(nodes, "_run_knowledge_tool", fake_run)
    await nodes.ingest_knowledge(
        {"task_id": "task-ingest", "user_input": "加入知识库 C:/Users/me/notes.md"}
    )
    assert captured == {
        "tool": "knowledge.ingest_file",
        "payload": {"path": "C:/Users/me/notes.md"},
    }


def test_explicit_markdown_file_import(knowledge_data_dir, tmp_path) -> None:
    source_file = tmp_path / "import-me.md"
    source_file.write_text("# Imported\n\nThis Markdown file contains enough content for an explicit knowledge import.", encoding="utf-8")
    result = ingest_file(str(source_file))
    assert result["status"] == "created"
    assert from_knowledge_relative(result["path"]).exists()


def test_change_storage_migrates_and_rebuilds(knowledge_data_dir, tmp_path) -> None:
    source = ingest_content(source_type="user", title="Relocated", content="Knowledge files must migrate when the configured storage directory changes.", canonical_uri="user://relocated", capture_method="test")
    target = tmp_path / "new-knowledge-root"
    result = change_storage_directory(str(target))
    assert result["migrated"] is True
    assert knowledge_root().resolve() == target.resolve()
    assert from_knowledge_relative(source["path"]).exists()
    assert result["rebuild"]["errors"] == 0


def test_change_storage_rejects_unrelated_nonempty_directory(knowledge_data_dir, tmp_path) -> None:
    target = tmp_path / "documents"
    target.mkdir()
    (target / "unrelated.txt").write_text("not a knowledge vault", encoding="utf-8")
    with pytest.raises(ValueError, match="非空"):
        change_storage_directory(str(target))


@pytest.mark.asyncio
async def test_ingest_compile_search_and_deduplicate(knowledge_data_dir, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        compiler,
        "get_settings",
        lambda: SimpleNamespace(openai_api_key=None, openai_base_url=None, openai_model="test"),
    )
    first = ingest_content(
        source_type="web",
        title="DeskPilot Knowledge Architecture",
        content="DeskPilot uses Markdown as the readable source and SQLite as its searchable control plane.",
        canonical_uri="https://example.com/deskpilot?utm_source=test",
        capture_method="test",
    )
    duplicate = ingest_content(
        source_type="web",
        title="DeskPilot Knowledge Architecture",
        content="DeskPilot uses Markdown as the readable source and SQLite as its searchable control plane.",
        canonical_uri="https://example.com/deskpilot",
        capture_method="test",
    )
    assert first["status"] == "created"
    assert duplicate["status"] == "already_exists"

    compiled = await compile_source(first["source_id"])
    assert len(compiled["committed"]) == 1
    note = get_note(compiled["committed"][0]["note_id"])
    assert note is not None
    assert note["status"] == "active"
    search_results = search_knowledge("DeskPilot Knowledge")
    assert search_results
    assert search_results[0]["updated_at"] == note["updated_at"]

    snapshot_path = from_knowledge_relative(first["path"])
    metadata, _ = parse_markdown_document(snapshot_path.read_text(encoding="utf-8"))
    assert metadata["id"] == first["source_id"]
    assert source_chunks(snapshot_path.read_text(encoding="utf-8"))["chunk-001"]

    client = TestClient(app)
    note_list = client.get("/knowledge/notes").json()
    assert note_list["total"] == 1
    assert client.get(f"/knowledge/notes/{note['id']}").json()["sections"]["Summary"]
    source_list = client.get("/knowledge/sources").json()
    assert source_list["total"] == 1
    assert source_list["items"][0]["snapshot_count"] == 1
    snapshot = client.get(f"/knowledge/snapshots/{first['snapshot_id']}").json()
    assert "Content" in snapshot["sections"]

    report = lint_knowledge()
    assert report["summary"]["errors"] == 0

    with connect() as connection:
        connection.execute("DELETE FROM knowledge_note_sources")
        connection.execute("DELETE FROM knowledge_relations")
        connection.execute("DELETE FROM knowledge_notes")
        connection.execute("DELETE FROM knowledge_snapshots")
        connection.execute("DELETE FROM knowledge_sources")
        connection.execute("DELETE FROM knowledge_fts")
    rebuilt = rebuild_index()
    assert rebuilt["errors"] == 0
    assert get_source(first["source_id"]) is not None
    assert get_note(note["id"]) is not None
    assert search_knowledge("DeskPilot Knowledge")


@pytest.mark.asyncio
async def test_source_update_creates_review_proposal(knowledge_data_dir, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        compiler,
        "get_settings",
        lambda: SimpleNamespace(openai_api_key=None, openai_base_url=None, openai_model="test"),
    )
    first = ingest_content(
        source_type="web",
        title="Versioned Source",
        content="The first version contains enough stable content to become a knowledge note.",
        canonical_uri="https://example.com/versioned",
        capture_method="test",
    )
    created = await compile_source(first["source_id"])
    note_id = created["committed"][0]["note_id"]

    updated = ingest_content(
        source_type="web",
        title="Versioned Source",
        content="The second version changes the source and should create a review proposal before updating the note.",
        canonical_uri="https://example.com/versioned",
        capture_method="test",
    )
    assert updated["status"] == "updated"
    assert get_note(note_id)["status"] == "stale"

    compilation = await compile_source(first["source_id"])
    assert compilation["pending"]
    proposal_id = compilation["pending"][0]["proposal_id"]
    assert list_proposals()[0]["id"] == proposal_id
    proposal_detail = TestClient(app).get(f"/knowledge/proposals/{proposal_id}")
    assert proposal_detail.status_code == 200
    assert proposal_detail.json()["payload"]["operation"]["operation"] == "update_note"
    assert proposal_detail.json()["target_note"]["id"] == note_id
    edit_response = TestClient(app).put(f"/knowledge/notes/{note_id}", json={"title": "Versioned Source", "entity_type": "note", "summary": "Manually protected old summary", "overview": "Old overview", "details": "Old details", "tags": []})
    assert edit_response.status_code == 200
    resolution = review_proposal(proposal_id, "accept")
    assert resolution["status"] == "accepted"
    assert get_note(note_id)["status"] == "active"
    accepted_detail = TestClient(app).get(f"/knowledge/notes/{note_id}").json()
    assert accepted_detail["sections"]["Summary"] == "Manually protected old summary"
    assert accepted_detail["sections"]["Overview"] == "Old overview"
    assert accepted_detail["sections"]["Details"] == "Old details"


@pytest.mark.asyncio
async def test_note_can_be_edited_and_deleted(knowledge_data_dir, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(compiler, "get_settings", lambda: SimpleNamespace(openai_api_key=None, openai_base_url=None, openai_model="test"))
    source = ingest_content(source_type="user", title="Editable Note", content="A stable source for testing manual knowledge page management.", canonical_uri="user://editable-note", capture_method="test")
    compiled = await compile_source(source["source_id"])
    note_id = compiled["committed"][0]["note_id"]
    client = TestClient(app)
    response = client.put(f"/knowledge/notes/{note_id}", json={"title": "Edited Product Method", "entity_type": "method", "summary": "Manual summary", "overview": "Manual overview", "details": "Manual details", "tags": ["product", "decision"]})
    assert response.status_code == 200
    edited = response.json()
    assert edited["entity_type"] == "method"
    assert edited["sections"]["Summary"] == "Manual summary"
    assert edited["frontmatter"]["manual_sections"] == ["Summary", "Overview", "Details"]
    edited_path = from_knowledge_relative(edited["markdown_path"])
    assert edited_path.exists()
    assert "Edited Product Method" in __import__("pathlib").Path(catalog_status()["files"]["Concept Index.md"]).read_text(encoding="utf-8")
    deleted = client.delete(f"/knowledge/notes/{note_id}")
    assert deleted.status_code == 200
    assert get_note(note_id) is None
    assert not edited_path.exists()
    with connect() as connection:
        assert connection.execute("SELECT COUNT(*) FROM knowledge_fts WHERE object_id=?", (note_id,)).fetchone()[0] == 0


@pytest.mark.asyncio
async def test_query_reads_matching_source_excerpt_when_note_omits_detail(knowledge_data_dir, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(compiler, "get_settings", lambda: SimpleNamespace(openai_api_key=None, openai_base_url=None, openai_model="test"))
    source = ingest_content(source_type="user", title="Compressed Paper", content="The paper contains the unique term hypergraph-retention-criterion in its detailed methodology.", canonical_uri="user://compressed-paper", capture_method="test")
    compiled = await compile_source(source["source_id"])
    note_id = compiled["committed"][0]["note_id"]
    response = TestClient(app).put(f"/knowledge/notes/{note_id}", json={"title": "Compressed Paper", "entity_type": "concept", "summary": "Short summary", "overview": "Short overview", "details": "", "tags": []})
    assert response.status_code == 200
    answer = await answer_knowledge("hypergraph-retention-criterion")
    assert answer["results"][0]["id"] == note_id
    excerpts = answer["evidence_bundle"]["notes"][0]["source_excerpts"]
    assert "hypergraph-retention-criterion" in excerpts[0]["content"]


@pytest.mark.asyncio
async def test_targeted_note_refresh_updates_selected_stale_note(knowledge_data_dir, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(compiler, "get_settings", lambda: SimpleNamespace(openai_api_key=None, openai_base_url=None, openai_model="test"))
    source = ingest_content(source_type="web", title="Original Title", content="The original source content is long enough to create the first page.", canonical_uri="https://example.com/targeted-refresh", capture_method="test")
    first = await compile_source(source["source_id"])
    stale_note_id = first["committed"][0]["note_id"]
    ingest_content(source_type="web", title="Renamed Title", content="The renamed source has changed content and should update the explicitly selected stale page.", canonical_uri="https://example.com/targeted-refresh", capture_method="test")
    unrelated = await compile_source(source["source_id"])
    assert unrelated["pending"] or unrelated["committed"]
    targeted = await compile_source(source["source_id"], stale_note_id)
    assert targeted["pending"][0]["proposal_id"]
    proposal = get_proposal(targeted["pending"][0]["proposal_id"])
    assert proposal["target_note_id"] == stale_note_id
    repeated = await compile_source(source["source_id"], stale_note_id)
    assert repeated["already_pending"] is True
    assert repeated["pending"][0]["proposal_id"] == targeted["pending"][0]["proposal_id"]


@pytest.mark.asyncio
async def test_targeted_model_proposal_keeps_relation_operations(knowledge_data_dir, monkeypatch: pytest.MonkeyPatch) -> None:
    target_id = "note_target"
    related_id = "note_related"
    payload = {
        "source_id": "source_1",
        "operations": [
            {
                "operation": "update_note",
                "target_note_id": target_id,
                "entity_type": "note",
                "title": "Target",
                "summary": "Updated summary",
                "details_markdown": "Updated details",
                "evidence": [{"source_id": "source_1", "snapshot_id": "snapshot_1", "anchor": "chunk-001", "reason": "test"}],
            },
            {
                "operation": "add_relation",
                "target_note_id": target_id,
                "relations": [{"relation_type": "related_to", "target_note_id": related_id, "confidence": 0.9}],
            },
            {
                "operation": "create_note",
                "title": "Unrequested extra page",
                "summary": "Should be filtered",
                "details_markdown": "Should be filtered",
                "evidence": [{"source_id": "source_1", "snapshot_id": "snapshot_1", "anchor": "chunk-001", "reason": "test"}],
            },
        ],
        "ignored_content": [],
    }

    class FakeCompletions:
        async def create(self, **kwargs):
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(payload)))])

    class FakeOpenAI:
        def __init__(self, **kwargs):
            self.chat = SimpleNamespace(completions=FakeCompletions())

    monkeypatch.setattr(compiler, "AsyncOpenAI", FakeOpenAI)
    monkeypatch.setattr(compiler, "get_settings", lambda: SimpleNamespace(openai_api_key="test", openai_base_url=None, openai_model="test"))
    monkeypatch.setattr(compiler, "get_note", lambda note_id: {"id": note_id, "title": "Existing"})
    proposal = await compiler._model_proposal(
        {"id": "source_1", "title": "Source", "sensitivity": "normal"},
        {"id": "snapshot_1"},
        {"chunk-001": "Evidence"},
        {"id": target_id, "title": "Target"},
    )
    assert proposal is not None
    assert [operation.operation for operation in proposal.operations] == ["update_note", "add_relation"]


@pytest.mark.asyncio
async def test_promote_answer_looks_up_existing_note_once_and_preserves_type(knowledge_data_dir, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(compiler, "get_settings", lambda: SimpleNamespace(openai_api_key=None, openai_base_url=None, openai_model="test"))
    source = ingest_content(source_type="user", title="Promotion Target", content="Traceable source evidence for the existing promotion target.", canonical_uri="user://promotion-target", capture_method="test")
    compiled = await compile_source(source["source_id"])
    note_id = compiled["committed"][0]["note_id"]
    response = TestClient(app).put(f"/knowledge/notes/{note_id}", json={"title": "Promotion Target", "entity_type": "method", "summary": "Manual summary", "overview": "Manual overview", "details": "Manual details", "tags": []})
    assert response.status_code == 200

    original_find = compiler.find_note_by_title
    calls = 0

    def counted_find(title: str):
        nonlocal calls
        calls += 1
        if calls > 1:
            raise AssertionError("find_note_by_title must be evaluated once")
        return original_find(title)

    monkeypatch.setattr(compiler, "find_note_by_title", counted_find)
    promoted = promote_answer("Promotion Target", "A promoted synthesis grounded in the existing note.", [note_id])
    assert calls == 1
    assert promoted["note_id"] == note_id
    assert get_note(note_id)["entity_type"] == "method"


def test_knowledge_settings_api_persists_values(knowledge_data_dir) -> None:
    client = TestClient(app)
    response = client.put(
        "/knowledge/settings",
        json={
            "enabled": True,
            "auto_compile": False,
            "review_updates": True,
            "allow_private_remote": False,
            "max_search_results": 12,
            "auto_create_notes": True,
            "purpose": "# Purpose\n\nOnly tested knowledge.",
        },
    )
    assert response.status_code == 200
    assert get_knowledge_settings().max_search_results == 12
    assert client.get("/knowledge/status").status_code == 200


def test_knowledge_status_does_not_rebuild_catalogs(knowledge_data_dir, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "backend.app.api.routes_knowledge.refresh_catalogs",
        lambda: pytest.fail("status reads must not rebuild catalogs"),
    )
    assert TestClient(app).get("/knowledge/status").status_code == 200


def test_whitespace_knowledge_root_uses_default(knowledge_data_dir) -> None:
    set_setting(KNOWLEDGE_ROOT_KEY, "   ")
    assert knowledge_root() == knowledge_data_dir / "knowledge"


def test_docx_and_image_import(knowledge_data_dir, tmp_path, monkeypatch: pytest.MonkeyPatch) -> None:
    from docx import Document
    from PIL import Image

    document_path = tmp_path / "research.docx"
    document = Document()
    document.add_heading("Research Notes", level=1)
    document.add_paragraph("This document contains enough structured material for the knowledge importer.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Topic"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Profiles"
    table.cell(1, 1).text = "Ready"
    document.save(document_path)
    docx_result = ingest_file(str(document_path))
    assert docx_result["status"] == "created"

    image_path = tmp_path / "scan.png"
    Image.new("RGB", (320, 120), "white").save(image_path)
    monkeypatch.setattr(
        "backend.app.rpa.ocr.ocr_image",
        lambda _: [{"text": "OCR imported knowledge content with sufficient detail.", "confidence": 0.98}],
    )
    image_result = ingest_file(str(image_path))
    assert image_result["status"] == "created"


def test_profile_isolation(knowledge_data_dir) -> None:
    first = ingest_content(
        source_type="user", title="Default Profile", content="Default profile knowledge content remains isolated from other profiles.",
        canonical_uri="user://default-profile", capture_method="test",
    )
    profile = create_profile("Research")
    switch_profile(profile["id"])
    assert get_source(first["source_id"]) is None
    second = ingest_content(
        source_type="user", title="Research Profile", content="Research profile knowledge content is independently visible and searchable.",
        canonical_uri="user://research-profile", capture_method="test",
    )
    assert get_source(second["source_id"]) is not None
    assert not search_knowledge("Default Profile")
    switch_profile("profile_default")
    assert get_source(first["source_id"]) is not None
    assert get_source(second["source_id"]) is None


def test_created_profile_matches_api_contract(knowledge_data_dir) -> None:
    profile = create_profile("Contract")
    assert profile["is_active"] is False
    assert profile["source_count"] == 0
    assert profile["note_count"] == 0


@pytest.mark.asyncio
async def test_web_update_check_detects_change(knowledge_data_dir, monkeypatch: pytest.MonkeyPatch) -> None:
    source = ingest_content(
        source_type="web", title="Watched Page", content="The original watched page has enough meaningful text for ingestion.",
        canonical_uri="https://example.com/watched", capture_method="test",
    )
    set_watch(source["source_id"], True, 60)
    async def fake_compile(source_id: str): return {"source_id": source_id, "committed": []}
    monkeypatch.setattr("backend.app.knowledge.web_monitor.compile_source", fake_compile)

    class FakeResponse:
        text = "<html><title>Watched Page</title><body><article>The changed watched page now contains substantially different knowledge content.</article></body></html>"
        status_code = 200
        url = "https://example.com/watched"
        def raise_for_status(self): return None

    class FakeClient:
        def __init__(self, **kwargs): pass
        async def __aenter__(self): return self
        async def __aexit__(self, *args): return None
        async def get(self, url): return FakeResponse()

    monkeypatch.setattr("backend.app.knowledge.web_monitor.httpx.AsyncClient", FakeClient)
    result = await check_web_source(source["source_id"])
    assert result["check_status"] == "changed"
    assert get_watch(source["source_id"])["last_changed_at"] is not None


@pytest.mark.asyncio
async def test_manual_web_check_enables_new_watch(knowledge_data_dir, monkeypatch: pytest.MonkeyPatch) -> None:
    source = ingest_content(
        source_type="web", title="Manual Check", content="Original content for a source that has no watch yet.",
        canonical_uri="https://example.com/manual-check", capture_method="test",
    )

    class FakeResponse:
        text = "<html><title>Manual Check</title><body><p>Original content for a source that has no watch yet.</p></body></html>"
        status_code = 200
        url = "https://example.com/manual-check"
        def raise_for_status(self): return None

    class FakeClient:
        def __init__(self, **kwargs): pass
        async def __aenter__(self): return self
        async def __aexit__(self, *args): return None
        async def get(self, url): return FakeResponse()

    monkeypatch.setattr("backend.app.knowledge.web_monitor.httpx.AsyncClient", FakeClient)
    await check_web_source(source["source_id"])
    assert get_watch(source["source_id"])["enabled"] == 1


def test_web_reingest_preserves_custom_watch_interval(knowledge_data_dir) -> None:
    source = ingest_content(source_type="web", title="Interval", content="Initial page content long enough for the knowledge source.", canonical_uri="https://example.com/interval", capture_method="test")
    set_watch(source["source_id"], True, 60)
    settings = get_knowledge_settings().model_copy(update={"auto_watch_web_sources": True, "web_update_interval_minutes": 1440})
    save_knowledge_settings(settings)
    ingest_content(source_type="web", title="Interval", content="Changed page content remains long enough and must preserve its custom schedule.", canonical_uri="https://example.com/interval", capture_method="test")
    assert get_watch(source["source_id"])["interval_minutes"] == 60


def test_chinese_search_tokenization() -> None:
    terms = search_terms("知识库支持中文全文检索")
    assert "知识库" in terms
    assert "全文" in terms


@pytest.mark.asyncio
async def test_obsidian_profile_export(knowledge_data_dir, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        compiler, "get_settings",
        lambda: SimpleNamespace(openai_api_key=None, openai_base_url=None, openai_model="test"),
    )
    source = ingest_content(
        source_type="user", title="Obsidian Integration",
        content="Obsidian integration exports profile notes and their traceable source material.",
        canonical_uri="user://obsidian-integration", capture_method="test",
    )
    await compile_source(source["source_id"])
    settings = get_knowledge_settings().model_copy(update={"obsidian_enabled": True})
    save_knowledge_settings(settings)
    result = export_obsidian_vault()
    vault = __import__("pathlib").Path(result["vault_path"])
    assert result["notes"] == 1
    assert (vault / "Home.md").exists()
    assert list((vault / "Notes").glob("*.md"))
    assert (vault / ".obsidian" / "app.json").exists()


@pytest.mark.asyncio
async def test_catalog_log_promotion_and_semantic_lint(knowledge_data_dir, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(compiler, "get_settings", lambda: SimpleNamespace(openai_api_key=None, openai_base_url=None, openai_model="test"))
    monkeypatch.setattr("backend.app.knowledge.lint.get_settings", lambda: SimpleNamespace(openai_api_key=None, openai_base_url=None, openai_model="test"))
    source = ingest_content(source_type="web", title="Compounding Wiki", content="A maintained wiki compounds synthesized knowledge and preserves source evidence over time.", canonical_uri="https://example.com/compounding-wiki", capture_method="test")
    compiled = await compile_source(source["source_id"])
    note_id = compiled["committed"][0]["note_id"]
    refresh_catalogs()
    files = catalog_status()["files"]
    for name in ["index.md", "Concept Index.md", "Dashboard.md", "log.md", "schema.md"]:
        assert __import__("pathlib").Path(files[name]).exists()
    concept_index = __import__("pathlib").Path(files["Concept Index.md"]).read_text(encoding="utf-8")
    assert "Compounding Wiki" in concept_index
    assert "## Note" in concept_index
    promoted = promote_answer("Why Wikis Compound", "The synthesis is retained and reused instead of reconstructed for every query.", [note_id])
    assert get_note(promoted["note_id"])["sources"]
    report = await semantic_lint_knowledge()
    assert report["scanned"]["notes"] == 2
    assert report["model_used"] is False
    log_content = __import__("pathlib").Path(files["log.md"]).read_text(encoding="utf-8")
    assert "ingest" in log_content and "compile" in log_content and "promote" in log_content

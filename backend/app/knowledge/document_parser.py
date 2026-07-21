from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class ParsedDocument:
    title: str
    content: str
    capture_method: str
    metadata: dict = field(default_factory=dict)


def parse_document(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        try:
            content = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("文件不是有效的 UTF-8 文本。") from exc
        return ParsedDocument(path.stem, content, "file_import", {"format": suffix[1:]})
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
        return _parse_image(path)
    raise ValueError("不支持的文件类型。")


def _parse_pdf(path: Path) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[str] = []
    page_methods: list[dict] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        meaningful_chars = sum(character.isalnum() or "\u4e00" <= character <= "\u9fff" for character in text)
        if meaningful_chars >= 40:
            pages.append(f"## Page {index}\n\n<!-- page:{index:03d} method:text -->\n{text}")
            page_methods.append({"page": index, "method": "text", "characters": len(text)})
            continue
        ocr_text, confidence, item_count = _ocr_pdf_page(path, index - 1)
        if ocr_text.strip():
            pages.append(f"## Page {index}\n\n<!-- page:{index:03d} method:ocr confidence:{confidence:.3f} -->\n{ocr_text}")
            page_methods.append({"page": index, "method": "ocr", "characters": len(ocr_text), "ocr_items": item_count, "ocr_average_confidence": confidence})
        elif text:
            pages.append(f"## Page {index}\n\n<!-- page:{index:03d} method:text-low-quality -->\n{text}")
            page_methods.append({"page": index, "method": "text-low-quality", "characters": len(text)})
    content = "\n\n".join(pages).strip()
    if not content:
        raise ValueError("PDF 文本层和逐页 OCR 均未识别到有效文字。")
    metadata = {
        "format": "pdf",
        "page_count": len(reader.pages),
        "document_metadata": {str(key): str(value) for key, value in (reader.metadata or {}).items()},
        "pages": page_methods,
        "text_pages": sum(item["method"] == "text" for item in page_methods),
        "ocr_pages": sum(item["method"] == "ocr" for item in page_methods),
    }
    title = str((reader.metadata or {}).get("/Title") or path.stem)
    method = "pdf_hybrid_ocr" if metadata["ocr_pages"] and metadata["text_pages"] else ("pdf_scanned_ocr" if metadata["ocr_pages"] else "pdf_text_extraction")
    return ParsedDocument(title, content, method, metadata)


def _ocr_pdf_page(path: Path, page_index: int) -> tuple[str, float, int]:
    import tempfile

    import pypdfium2 as pdfium

    from backend.app.rpa.ocr import ocr_image

    document = pdfium.PdfDocument(str(path))
    try:
        page = document[page_index]
        bitmap = page.render(scale=2.5, rotation=0)
        image = bitmap.to_pil()
        with tempfile.TemporaryDirectory(prefix="deskpilot-pdf-ocr-") as temporary:
            image_path = Path(temporary) / f"page-{page_index + 1}.png"
            image.save(image_path, format="PNG")
            items = ocr_image(image_path)
    finally:
        document.close()
    accepted = [item for item in items if float(item.get("confidence", 0)) >= 0.45 and str(item.get("text", "")).strip()]
    if not accepted:
        return "", 0.0, 0
    confidence = sum(float(item.get("confidence", 0)) for item in accepted) / len(accepted)
    return "\n".join(str(item["text"]).strip() for item in accepted), confidence, len(accepted)


def _parse_docx(path: Path) -> ParsedDocument:
    from docx import Document

    document = Document(str(path))
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else ""
        if style.startswith("Heading"):
            level = min(3, max(1, int(style.split()[-1]) if style.split()[-1].isdigit() else 2))
            blocks.append(f"{'#' * level} {text}")
        else:
            blocks.append(text)
    for table_index, table in enumerate(document.tables, start=1):
        rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        normalized = [row + [""] * (width - len(row)) for row in rows]
        lines = [f"## Table {table_index}", "", "| " + " | ".join(normalized[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
        lines.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
        blocks.append("\n".join(lines))
    content = "\n\n".join(blocks).strip()
    if not content:
        raise ValueError("DOCX 没有可提取的正文或表格。")
    core = document.core_properties
    return ParsedDocument(
        core.title or path.stem,
        content,
        "docx_structured_extraction",
        {
            "format": "docx",
            "paragraph_count": len(document.paragraphs),
            "table_count": len(document.tables),
            "author": core.author or "",
        },
    )


def _parse_image(path: Path) -> ParsedDocument:
    from PIL import Image

    from backend.app.rpa.ocr import ocr_image

    with Image.open(path) as image:
        dimensions = {"width": image.width, "height": image.height, "mode": image.mode}
    items = ocr_image(path)
    accepted = [item for item in items if float(item.get("confidence", 0)) >= 0.45 and str(item.get("text", "")).strip()]
    content = "\n".join(str(item["text"]).strip() for item in accepted)
    if not content.strip():
        raise ValueError("图片 OCR 没有识别到有效文字。")
    average = sum(float(item["confidence"]) for item in accepted) / len(accepted)
    return ParsedDocument(
        path.stem,
        content,
        "paddleocr_image",
        {"format": path.suffix.lower()[1:], "ocr_items": len(accepted), "ocr_average_confidence": average, **dimensions},
    )
